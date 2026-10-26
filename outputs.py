from __future__ import annotations

from io import BytesIO
from typing import Any

import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape, portrait
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from engineering import METRIC_DESCRIPTIONS, find_curve, find_model


def _rev_text(v: Any) -> str:
    s = str(v or "A").strip().upper()
    return s or "A"


def build_excel(project: dict[str, Any], nozzle: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DESIGN DATA SHEET"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.28
    ws.page_margins.right = 0.28
    ws.page_margins.top = 0.32
    ws.page_margins.bottom = 0.32

    widths = {"A": 2.5, "B": 5, "C": 3, "D": 22, "E": 4, "F": 5, "G": 3, "H": 3,
              "I": 9, "J": 9, "K": 9, "L": 9, "M": 9, "N": 4}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w

    thin = Side(style="thin", color="000000")
    value_border = Border(bottom=thin)
    title_font = Font(name="Malgun Gothic", size=16, bold=True)
    bold = Font(name="Malgun Gothic", size=10, bold=True)
    normal = Font(name="Malgun Gothic", size=10)

    ws.merge_cells("B1:M1")
    ws["B1"] = "ATOMIZING NOZZLE DESIGN DATA SHEET"
    ws["B1"].font = title_font
    ws["B1"].alignment = Alignment(horizontal="center")
    ws["N1"] = "Rev." + _rev_text(nozzle.get("data_sheet_revision"))
    ws["N1"].alignment = Alignment(horizontal="right")

    titles = {"RA": "R.A Injection Nozzle", "DI": "Direct Injection Nozzle", "AMMONIA": "Ammonia Injection Nozzle"}
    ws.merge_cells("B3:N3")
    ws["B3"] = "1. " + titles.get(nozzle.get("injection_type"), "Injection Nozzle") + (" / " + nozzle.get("selected_model", "") if nozzle.get("selected_model") else "")
    ws["B3"].font = bold

    qty = max(1, int(nozzle.get("qty") or 1))
    fluid = nozzle.get("fluid", "")
    if nozzle.get("concentration") and fluid != "Water":
        fluid += " " + str(nozzle.get("concentration"))
    cap = f"Min {nozzle.get('min_flow_lmin',0):g} / Normal {nozzle.get('normal_flow_lmin',0):g} / Max {nozzle.get('max_flow_lmin',0):g} L/min / nozzle"
    man = nozzle.get("manual", {})

    rows = [
        ("1)", "Quantity", f"Total ({qty} sets/Unit)"),
        ("2)", "Model/Spraying Type", " / ".join(x for x in [nozzle.get("selected_model", ""), man.get("nozzle_type", ""), man.get("run_type", "")] if x)),
        ("3)", "Capacity", cap),
        ("4)", "Liquid", ""),
        ("", "a. Liquid", fluid),
        ("", "- Droplet Size", nozzle.get("droplet_size", "")),
        ("", "- Spraying Degree", f"{nozzle.get('spray_angle_deg',0):g}°"),
        ("", "- Spraying Type", man.get("spraying_type", "")),
        ("", "b. Atomizing Air", "Plant Air"),
        ("", "- Pressure", f"{nozzle.get('air_pressure_barg',0):g} barg" if nozzle.get("air_pressure_barg") else ""),
        ("", "- Flow rate", nozzle.get("air_flow_nm3h", "")),
        ("5)", "Liquid Pressure", nozzle.get("liquid_pressure_barg", "")),
        ("6)", "Dimension", man.get("dimension", "")),
        ("7)", "Material", ""),
        ("", "- Nozzle Tip", man.get("tip_material", "")),
        ("", "- Nozzle Pipe", man.get("pipe_material", "")),
        ("8)", "Weight", ""),
        ("", "- Empty", man.get("empty_weight", "")),
        ("", "- Operating", man.get("operating_weight", "")),
        ("9)", "Connection", ""),
        ("", "- Atomizing air", man.get("air_connection", "")),
        ("", "- Liquid", man.get("liquid_connection", "")),
        ("10)", "Unit Number", man.get("unit_number", "")),
        ("11)", "Operation Temperature", ""),
        ("", "- HRSG Design", man.get("design_temperature", "")),
        ("", "- HRSG Operation", man.get("operating_temperature", "")),
    ]

    r = 5
    for no, label, value in rows:
        ws.cell(r, 2, no).font = bold if no else normal
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        ws.cell(r, 4, label).font = bold if no or label in ("a. Liquid", "b. Atomizing Air") else normal
        ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=13)
        ws.cell(r, 9, value)
        ws.cell(r, 9).font = normal
        ws.cell(r, 9).alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True)
        for c in range(9, 14):
            ws.cell(r, c).border = value_border
        r += 1

    if nozzle.get("additional_specs"):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
        ws.cell(r, 2, "Additional Specification").font = bold
        r += 1
        for a in nozzle["additional_specs"]:
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
            ws.cell(r, 4, str(a.get("label", "")))
            ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=13)
            value = str(a.get("value", "")) + ((" " + str(a.get("unit"))) if a.get("unit") else "")
            ws.cell(r, 9, value)
            for c in range(9, 14): ws.cell(r, c).border = value_border
            r += 1
    if man.get("remark"):
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        ws.cell(r, 4, "Remark / Note").font = bold
        ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=13)
        ws.cell(r, 9, man["remark"])
        for c in range(9, 14): ws.cell(r, c).border = value_border

    buf = BytesIO(); wb.save(buf); return buf.getvalue()


def _graph_png(model: dict[str, Any], curve: dict[str, Any], nozzle: dict[str, Any], flow_unit: str, show_points: bool) -> bytes:
    data = curve.get("data", [])
    factor = 1/60 if flow_unit == "L/min" else 1
    x = [p.get("flow", 0) * factor for p in data]
    fig, ax1 = plt.subplots(figsize=(10.5, 5.0), dpi=150)
    ax_air = ax1.twinx()
    ax2 = ax1.twinx()
    ax_air.spines["right"].set_visible(False)
    ax_air.spines["left"].set_position(("outward", 42))
    ax_air.spines["left"].set_visible(True)
    ax_air.yaxis.set_label_position("left")
    ax_air.yaxis.set_ticks_position("left")
    if any("liquidPressure" in p for p in data): ax1.plot(x, [p.get("liquidPressure") for p in data], marker="o", label="Liquid Pressure (barg)")
    if any("airRate" in p for p in data): ax_air.plot(x, [p.get("airRate") for p in data], marker="o", label="Air Flow (Nm3/hr)")
    for key, label in [("d32","D32/SMD"),("dv05","Dv0.5"),("dv09","Dv0.9"),("dmax","Dmax"),("x90","X90"),("x99","X99")]:
        if any(key in p for p in data): ax2.plot(x, [p.get(key) for p in data], marker="o", label=label)
    if show_points:
        for i, dp in enumerate(nozzle.get("design_points", []), 1):
            fx = float(dp.get("flow_lhr", 0)) * factor
            ax1.axvline(fx, linestyle="--", linewidth=1)
            ax1.text(fx, ax1.get_ylim()[1], f"P{i}", ha="center", va="top", fontsize=8)
    ax1.set_title(f"{model.get('model')} · Atomizing Air {curve.get('air')} barg")
    ax1.set_xlabel(f"Liquid Flow ({flow_unit})")
    ax1.set_ylabel("Liquid Pressure (barg)")
    ax_air.set_ylabel("Air Flow (Nm3/hr)")
    ax2.set_ylabel("Particle Size (micrometres)")
    ax1.grid(True, alpha=.2)
    h1,l1=ax1.get_legend_handles_labels(); ha,la=ax_air.get_legend_handles_labels(); h2,l2=ax2.get_legend_handles_labels()
    if h1 or ha or h2: ax1.legend(h1+ha+h2,l1+la+l2,loc="best",fontsize=8)
    fig.tight_layout()
    b=BytesIO(); fig.savefig(b,format="png",bbox_inches="tight"); plt.close(fig); return b.getvalue()


def build_performance_pdf(project: dict[str, Any], nozzle: dict[str, Any], perf_db: list[dict[str, Any]], flow_unit: str = "L/min") -> bytes:
    model = find_model(perf_db, nozzle.get("selected_model", ""))
    curve = find_curve(model, nozzle.get("selected_air_barg"))
    if not model or not curve:
        raise ValueError("Select a digital model and air pressure first.")
    opt = nozzle.get("pdf_options", {})
    page = landscape(A4) if opt.get("orientation", "landscape") == "landscape" else portrait(A4)
    buf = BytesIO()
    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=26, rightMargin=26, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    title = styles["Title"]; title.fontName = "Helvetica-Bold"; title.fontSize = 18
    body = styles["BodyText"]; body.fontName = "HYSMyeongJo-Medium"; body.fontSize = 8.5; body.leading = 11
    story = [Paragraph("ATOMIZING NOZZLE PERFORMANCE CURVE", title), Spacer(1, 5)]
    story.append(Paragraph(f"Revision: {_rev_text(nozzle.get('performance_curve_revision'))}", body))
    if opt.get("project_info", True):
        info = [
            ["Project", project.get("project_name", ""), "Customer", project.get("customer", "")],
            ["Project No.", project.get("project_no", ""), "Nozzle", nozzle.get("name", "")],
            ["Model", model.get("model", ""), "Air Pressure", f"{curve.get('air')} barg"],
        ]
        t=Table(info,colWidths=[70,170,75,170]); t.setStyle(TableStyle([
            ("FONTNAME",(0,0),(-1,-1),"HYSMyeongJo-Medium"),("FONTSIZE",(0,0),(-1,-1),8),
            ("BOTTOMPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),4),
            ("LINEBELOW",(1,0),(1,-1),.4,colors.grey),("LINEBELOW",(3,0),(3,-1),.4,colors.grey),
        ])); story += [t, Spacer(1,8)]

    png = _graph_png(model, curve, nozzle, flow_unit, opt.get("graph_design_points", True))
    image = Image(BytesIO(png))
    maxw = page[0]-60; maxh = 205 if page[0] > page[1] else 400
    scale = min(maxw/image.imageWidth, maxh/image.imageHeight)
    image.drawWidth=image.imageWidth*scale; image.drawHeight=image.imageHeight*scale
    story += [image, Spacer(1,3)]

    if opt.get("point_table", True) and nozzle.get("design_points"):
        rows=[["Point",f"Flow ({flow_unit})","Liquid P. (barg)","Air (Nm3/hr)","D32","Dv0.9","Dmax","X90","X99"]]
        for i,p in enumerate(nozzle["design_points"],1):
            f=float(p.get("flow_lhr",0))/(60 if flow_unit=="L/min" else 1)
            rows.append([f"P{i}",f"{f:.3g}",_fmt(p.get("liquidPressure")),_fmt(p.get("airRate")),_fmt(p.get("d32") or p.get("smd")),_fmt(p.get("dv09")),_fmt(p.get("dmax")),_fmt(p.get("x90")),_fmt(p.get("x99"))])
        t=Table(rows,repeatRows=1); t.setStyle(TableStyle([
            ("FONTNAME",(0,0),(-1,-1),"Helvetica"),("FONTSIZE",(0,0),(-1,-1),7),
            ("BACKGROUND",(0,0),(-1,0),colors.whitesmoke),("GRID",(0,0),(-1,-1),.3,colors.lightgrey),
            ("ALIGN",(1,1),(-1,-1),"RIGHT")
        ])); story += [t,Spacer(1,3)]

    if opt.get("metric_descriptions", True):
        available = {"flow"}
        for p in curve.get("data",[]):
            if "liquidPressure" in p: available.add("liquidPressure")
            if "airRate" in p: available.add("airRate")
            for k in ("d32","dv05","dv09","dmax","x90","x99"):
                if k in p: available.add(k)
        selected=opt.get("metrics",{})
        desc=[]
        for k,(label,text) in METRIC_DESCRIPTIONS.items():
            if k in available and selected.get(k,True): desc.append([label,text])
        if desc:
            story.append(Paragraph("GRAPH METRIC DESCRIPTIONS", styles["Heading3"]))
            t=Table(desc,colWidths=[110,page[0]-60-110]); t.setStyle(TableStyle([
                ("FONTNAME",(0,0),(-1,-1),"Helvetica"),("FONTSIZE",(0,0),(-1,-1),6.8),
                ("VALIGN",(0,0),(-1,-1),"TOP"),("BOTTOMPADDING",(0,0),(-1,-1),3),
                ("FONTNAME",(0,0),(0,-1),"Helvetica-Bold")
            ])); story.append(t)

    story += [Paragraph(f"Source: {model.get('source','')} · Display flow unit: {flow_unit}", body)]
    doc.build(story); return buf.getvalue()


def _fmt(v: Any) -> str:
    if v is None or v == "": return "-"
    try: return f"{float(v):.3g}"
    except Exception: return str(v)
